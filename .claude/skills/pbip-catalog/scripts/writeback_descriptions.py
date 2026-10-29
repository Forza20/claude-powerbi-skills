#!/usr/bin/env python3
"""
writeback_descriptions.py - Write approved catalog descriptions back to TMDL as `///` lines.

Usage:
    python3 writeback_descriptions.py <pbip-folder> <catalog.xlsx> [--apply]

Only rows on the sheets Tabeller / Kolonner / Measures with Status = "Godkendt" and a
non-empty Beskrivelse are written. Without --apply the script only shows what would
change (dry run). Only `///` description lines are touched - never DAX, M or properties.
Close the project in Power BI Desktop before --apply, or Desktop may overwrite the files.
"""

import argparse
import difflib
import re
import sys
from pathlib import Path


def split_decl(line):
    """`\t<kind> <name> [= ...]` -> (indent, kind, name) or None."""
    m = re.match(r'^([\t ]*)(table|column|measure)\s+(.*)$', line.rstrip('\r\n'))
    if not m:
        return None
    indent, kind, rest = m.groups()
    rest = rest.strip()
    if rest.startswith("'"):
        i, buf = 1, []
        while i < len(rest):
            if rest[i] == "'":
                if i + 1 < len(rest) and rest[i + 1] == "'":
                    buf.append("'")
                    i += 2
                    continue
                break
            buf.append(rest[i])
            i += 1
        name = ''.join(buf)
    else:
        name = re.match(r'([^\s=]+)', rest).group(1)
    return indent, kind, name


def parse_key(key, sheet):
    if sheet == 'Tabeller':
        return key, None
    m = re.match(r'^(.+?)\[(.+)\]$', key)
    return (m.group(1), m.group(2)) if m else (None, None)


def load_approved(xlsx):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    out = []
    for sheet, kind in (('Tabeller', 'table'), ('Kolonner', 'column'), ('Measures', 'measure')):
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        head = [str(h or '') for h in rows[0]] if rows else []
        if not {'Nøgle', 'Beskrivelse', 'Status'} <= set(head):
            continue
        ix = {h: i for i, h in enumerate(head)}
        for r in rows[1:]:
            key, desc, status = r[ix['Nøgle']], r[ix['Beskrivelse']], r[ix['Status']]
            if key and desc and str(status).strip() == 'Godkendt':
                table, name = parse_key(str(key), sheet)
                if table:
                    out.append({'kind': kind, 'table': table, 'name': name or table,
                                'description': str(desc).strip()})
    return out


def find_table_file(tables_dir, table):
    for f in sorted(tables_dir.glob('*.tmdl')):
        for line in f.read_text(encoding='utf-8').splitlines():
            d = split_decl(line)
            if d and d[1] == 'table':
                if d[2] == table:
                    return f
                break
    return None


def apply_to_text(text, items):
    nl = '\r\n' if '\r\n' in text else '\n'
    lines = text.split(nl)
    changed = []
    for it in items:
        target = None
        for i, line in enumerate(lines):
            d = split_decl(line)
            if d and d[1] == it['kind'] and d[2] == it['name']:
                target = i
                break
        if target is None:
            changed.append((it, 'ikke fundet'))
            continue
        indent = split_decl(lines[target])[0]
        start = target
        while start > 0 and lines[start - 1].strip().startswith('///'):
            start -= 1
        new_desc = [f'{indent}/// {l.strip()}'.rstrip() for l in it['description'].splitlines() if l.strip()]
        old_desc = lines[start:target]
        if [l.strip() for l in old_desc] == [l.strip() for l in new_desc]:
            changed.append((it, 'uændret'))
            continue
        lines[start:target] = new_desc
        changed.append((it, 'opdateret' if old_desc else 'tilføjet'))
    return nl.join(lines), changed


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('pbip')
    ap.add_argument('catalog')
    ap.add_argument('--apply', action='store_true')
    a = ap.parse_args()

    sm = next(Path(a.pbip).glob('*.SemanticModel'), None)
    tables_dir = sm / 'definition' / 'tables' if sm else None
    if not tables_dir or not tables_dir.is_dir():
        sys.exit('Fandt ikke <navn>.SemanticModel/definition/tables (kun TMDL-format understøttes).')

    approved = load_approved(a.catalog)
    if not approved:
        print('Ingen rækker med Status = Godkendt. Intet at skrive.')
        return
    by_file, missing = {}, []
    for it in approved:
        f = find_table_file(tables_dir, it['table'])
        if f:
            by_file.setdefault(f, []).append(it)
        else:
            missing.append(it)

    summary = {'tilføjet': 0, 'opdateret': 0, 'uændret': 0, 'ikke fundet': len(missing)}
    for f, items in by_file.items():
        old = f.read_text(encoding='utf-8')
        new, changes = apply_to_text(old, items)
        for it, what in changes:
            summary[what] += 1
            if what == 'ikke fundet':
                missing.append(it)
        if new != old:
            diff = difflib.unified_diff(old.splitlines(), new.splitlines(),
                                        fromfile=f.name, tofile=f.name, lineterm='', n=1)
            print('\n'.join(diff))
            if a.apply:
                f.write_text(new, encoding='utf-8', newline='')
    for it in missing:
        print(f"! Ikke fundet i TMDL: {it['kind']} {it['table']}[{it['name']}]")
    mode = 'SKREVET' if a.apply else 'TØRKØRSEL (brug --apply for at skrive)'
    print(f"\n{mode}: {summary['tilføjet']} tilføjet, {summary['opdateret']} opdateret, "
          f"{summary['uændret']} uændret, {summary['ikke fundet']} ikke fundet.")


if __name__ == '__main__':
    main()
