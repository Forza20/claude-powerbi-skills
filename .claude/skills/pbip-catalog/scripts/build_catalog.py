#!/usr/bin/env python3
"""
build_catalog.py - Data catalog, lineage and dbdiagram.io model for a PBIP project.

Usage:
    python3 build_catalog.py <pbip-folder> [--out DIR] [--layout-config FILE]
                             [--previous catalog.xlsx] [--ai-descriptions FILE]
                             [--fence azdo|github] [--parser parse_pbip.py]

Writes to --out (default <pbip-folder>/docs):
    catalog.xlsx               Excel catalog (edit Beskrivelse/Ejer/Følsomhed/Opstrøms kilde/Status)
    lineage.md                 Mermaid lineage: overview, per dashboard slot, per column
    model.dbml                 DBML for dbdiagram.io / dbdocs.io
    catalog.json               Machine-readable catalog (pipeline handoff / diff)
    missing_descriptions.json  Items without a description + context, for AI drafts

Description precedence: earlier Excel edits > TMDL `///` > AI draft (--ai-descriptions).
"""

import argparse
import datetime
import importlib.util
import json
import os
import re
import shutil
import sys
from pathlib import Path

LAYOUT_PAGE_ID = '8d1e4c7a2b5f9e3d6c10'   # fixed page id written by pbip-layout
EDITABLE = ['Beskrivelse', 'Ejer', 'Følsomhed', 'Opstrøms kilde', 'Status']
STATUSES = ['Mangler', 'AI-forslag', 'Fra model', 'Redigeret', 'Godkendt']
PII_DEFAULT = 'Mulig PII - vurdér'


# ---------------------------------------------------------------------------
# Parser loading (shared with pbip-governance)
# ---------------------------------------------------------------------------

def load_parser(explicit=None):
    sys.dont_write_bytecode = True  # do not leave __pycache__ in the skill folders
    here = Path(__file__).resolve()
    candidates = [explicit, os.environ.get('PBIP_PARSER')]
    candidates += [str(p / 'skills' / 'pbip-governance' / 'scripts' / 'parse_pbip.py') for p in here.parents]
    candidates += [str(p / 'pbip-governance' / 'scripts' / 'parse_pbip.py') for p in here.parents]
    candidates += [
        str(Path.home() / '.claude' / 'skills' / 'pbip-governance' / 'scripts' / 'parse_pbip.py'),
        '/mnt/skills/user/pbip-governance/scripts/parse_pbip.py',
    ]
    for c in candidates:
        if c and Path(c).is_file():
            spec = importlib.util.spec_from_file_location('parse_pbip', c)
            mod = importlib.util.module_from_spec(spec)
            spec.loader.exec_module(mod)
            if not hasattr(mod, 'enrich_model'):
                continue  # old parser version without lineage support
            return mod, c
    sys.exit('Fandt ikke pbip-governance/scripts/parse_pbip.py (med lineage-understøttelse). '
             'Angiv stien med --parser eller miljøvariablen PBIP_PARSER.')


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def fkey(table, name):
    return f'{table}[{name}]'


def parse_fkey(s):
    m = re.match(r"^'?(.+?)'?\[(.+)\]$", (s or '').strip())
    return (m.group(1), m.group(2)) if m else (None, None)


def short(s, n=80):
    s = re.sub(r'\s+', ' ', str(s or '')).strip()
    return s if len(s) <= n else s[:n - 1] + '…'


def classify_table(t, rels):
    n = t['name'].lower()
    if t['measures'] and all(c['isHidden'] for c in t['columns']):
        return 'Measures'
    if t.get('dataCategory', '').lower() == 'time' or re.match(r'^(dim_)?(date|dato|kalender|calendar)', n):
        return 'Dato'
    if n.startswith(('fact_', 'fct_', 'f_', 'fakta')):
        return 'Fakta'
    if n.startswith(('dim_', 'd_')):
        return 'Dimension'
    if n.startswith(('bridge', 'bro_')):
        return 'Bro'
    many = any(r['from']['table'] == t['name'] for r in rels)
    one = any(r['to']['table'] == t['name'] for r in rels)
    if many and not one:
        return 'Fakta'
    if one:
        return 'Dimension'
    return 'Hjælpe'


def source_text(src):
    if not src:
        return ''
    if src.get('kind') in ('Other', 'Unknown'):
        return ''
    return f"{src['kind']}: {src.get('details', '')}"


def origin_text(origin):
    if not origin:
        return ''
    k = origin.get('kind')
    if k == 'source':
        obj = origin.get('object') or origin.get('query', '')
        return f"{obj}.{origin.get('column', '')}"
    if k == 'computed':
        step = origin.get('step', '')
        return f"Beregnet i {origin.get('query', '')}" + (f" (trin {step})" if step and step != '__expr__' else '')
    if k == 'join':
        return f"{origin.get('query', '')}[{origin.get('column', '')}]"
    return ''


def lineage_text(events):
    parts = []
    for e in events or []:
        if e['action'] == 'Datatype sat':
            continue
        d = short(e.get('detail', ''), 60)
        parts.append(f"{e['query']}: {e['action']}" + (f" ({d})" if d else ''))
    return ' → '.join(parts)


# ---------------------------------------------------------------------------
# Dashboard slots (pbip-layout)
# ---------------------------------------------------------------------------

def slots_from_config(cfg):
    out = []

    def add(slot, field, role, label=''):
        if field:
            out.append({'slot': slot, 'field': field, 'role': role, 'label': label})
    for i, k in enumerate(cfg.get('kpis', []), 1):
        add(f'KPI {i}', k.get('measure'), 'Værdi', k.get('label', ''))
    tr = cfg.get('trend', {})
    add('Trend', tr.get('measure'), 'Værdi', tr.get('title', ''))
    add('Trend', tr.get('monthColumn'), 'Akse', tr.get('title', ''))
    for i, m in enumerate(cfg.get('monthly', []), 1):
        for key, role in (('sparkMeasure', 'Sparkline'), ('currentMeasure', 'Seneste måned'),
                          ('labelMeasure', 'Forrige måned-tekst'), ('colorMeasure', 'Farve')):
            add(f'Månedskort {i}', m.get(key), role, m.get('title', ''))
    b = cfg.get('bar', {})
    add('Søjler pr. kategori', b.get('measure'), 'Værdi', b.get('title', ''))
    add('Søjler pr. kategori', b.get('categoryColumn'), 'Kategori', b.get('title', ''))
    tb = cfg.get('table', {})
    add('Top 10-tabel', tb.get('rowColumn'), 'Række', tb.get('title', ''))
    for c in tb.get('columns', []):
        add('Top 10-tabel', c.get('measure'), 'Kolonne', tb.get('title', ''))
    for key, role in (('sortMeasure', 'Sortering'), ('dataBarMeasure', 'Data bars'),
                      ('gradientMeasure', 'Farvegradient')):
        add('Top 10-tabel', tb.get(key), role, tb.get('title', ''))
    add('Top 10-tabel', (tb.get('topN') or {}).get('orderByColumn'), 'TopN-filter', tb.get('title', ''))
    for i, h in enumerate(cfg.get('highlights', []), 1):
        add(f'Fremhævning {i}', h.get('measure'), 'Værdi', h.get('header', ''))
    return out


ROLE_DA = {'Y': 'Værdi', 'Values': 'Værdi', 'Data': 'Værdi', 'Category': 'Kategori', 'Rows': 'Række',
           'Columns': 'Kolonne', 'Tooltips': 'Tooltip', 'Legend': 'Forklaring', 'Series': 'Serie',
           'X': 'Akse', 'Size': 'Størrelse'}


def slots_from_page(page):
    """Fallback without layout-config.json: recognise the pbip-layout page by positions (layout.md)."""
    out = []

    def bucket(x, starts):
        return min(range(len(starts)), key=lambda i: abs(starts[i] - x)) + 1
    for v in page['visuals']:
        p = v.get('position', {})
        x, y, typ = p.get('x', 0), p.get('y', 0), v.get('type')
        slot = None
        if typ == 'cardVisual' and abs(y - 20) < 10:
            slot = f'KPI {bucket(x, [376, 764, 1152, 1540])}'
        elif typ == 'lineChart':
            slot = 'Trend'
        elif y >= 760 and x < 1060 and typ in ('areaChart', 'cardVisual'):
            slot = f'Månedskort {bucket(x, [96, 422, 748])}'
        elif typ == 'clusteredBarChart':
            slot = 'Søjler pr. kategori'
        elif typ == 'tableEx':
            slot = 'Top 10-tabel'
        elif typ == 'cardVisual' and y >= 930:
            slot = f'Fremhævning {bucket(x, [1072, 1500])}'
        if not slot:
            continue
        for r in v.get('fieldRefs', []):
            out.append({'slot': slot, 'field': fkey(r['table'], r['name']),
                        'role': ROLE_DA.get(r['role'], r['role']), 'label': v.get('title', '')})
    return out


# ---------------------------------------------------------------------------
# Previous Excel (round-trip)
# ---------------------------------------------------------------------------

def read_previous(path):
    prev = {}
    if not path or not Path(path).is_file():
        return prev
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet in ('Tabeller', 'Kolonner', 'Measures'):
        if sheet not in wb.sheetnames:
            continue
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            continue
        head = [str(h or '') for h in rows[0]]
        if 'Nøgle' not in head:
            continue
        for r in rows[1:]:
            d = {h: (r[i] if i < len(r) else None) for i, h in enumerate(head)}
            if d.get('Nøgle'):
                d['_sheet'] = sheet
                prev[str(d['Nøgle'])] = d
    return prev


def merge_editable(key, tmdl_desc, ai, prev, defaults=None):
    """Apply the description precedence. Returns dict with the editable columns."""
    p = prev.get(key) or {}
    out = {k: (p.get(k) or '') for k in EDITABLE}
    for k, v in (defaults or {}).items():
        if not out.get(k):
            out[k] = v
    prev_desc = str(p.get('Beskrivelse') or '').strip()
    prev_status = str(p.get('Status') or '').strip()
    tmdl_desc = (tmdl_desc or '').strip()
    if prev_desc and prev_status != 'Fra model':
        out['Beskrivelse'], out['Status'] = prev_desc, prev_status or 'Redigeret'
        if prev_status == 'AI-forslag' and ai.get(key) and ai[key].strip() != prev_desc:
            pass  # keep what is in Excel; a new AI draft never overwrites it
        src = 'Excel'
    elif tmdl_desc:
        out['Beskrivelse'], out['Status'], src = tmdl_desc, 'Fra model', 'TMDL'
    elif ai.get(key):
        out['Beskrivelse'], out['Status'], src = ai[key].strip(), 'AI-forslag', 'AI'
    else:
        out['Beskrivelse'], out['Status'], src = '', 'Mangler', ''
    if out['Status'] not in STATUSES:
        out['Status'] = 'Redigeret'
    out['Beskrivelseskilde'] = src
    return out


# ---------------------------------------------------------------------------
# Catalog assembly
# ---------------------------------------------------------------------------

def build(parsed, slots, prev, ai, pii_reason):
    model, report = parsed['model'], parsed['report']
    rels = model['relationships']
    tables = model['tables']
    measure_table = {m['name']: t['name'] for t in tables for m in t['measures']}
    mdeps = {m['name']: m.get('dependencies', {}) for t in tables for m in t['measures']}

    # transitive closure per measure
    def closure(mname, seen=None):
        seen = seen if seen is not None else set()
        cols, tabs = set(), set()
        for d in mdeps.get(mname, {}).get('measures', []):
            if d not in seen:
                seen.add(d)
                c2, t2, _ = closure(d, seen)
                cols |= c2
                tabs |= t2
        for c in mdeps.get(mname, {}).get('columns', []):
            cols.add(fkey(c['table'], c['column']))
        tabs |= set(mdeps.get(mname, {}).get('tables', []))
        return cols, tabs, seen

    used_by = {}   # measure -> measures that use it directly
    for name, d in mdeps.items():
        for dep in d.get('measures', []):
            used_by.setdefault(dep, set()).add(name)

    # dashboard slot tagging (direct + indirect)
    direct, indirect = {}, {}
    for s in slots:
        direct.setdefault(s['field'], []).append(f"{s['slot']} ({s['role']})")
        t, n = parse_fkey(s['field'])
        if n in mdeps:
            cols, tabs, ms = closure(n)
            for m in ms:
                indirect.setdefault(fkey(measure_table[m], m), set()).add(s['slot'])
            for c in cols:
                indirect.setdefault(c, set()).add(s['slot'])
            for tb in tabs:
                indirect.setdefault(tb, set()).add(s['slot'])
        if t:
            indirect.setdefault(t, set()).add(s['slot'])

    # report usage (all pages)
    usage = {}
    for page in report.get('pages', []):
        for v in page.get('visuals', []):
            refs = v.get('fieldRefs')
            if refs is None:
                refs = []
                for f in v.get('fields', []):
                    if '.' in f:
                        a, b = f.split('.', 1)
                        refs.append({'table': a, 'name': b})
            label = f"{page.get('displayName', '')} › {v.get('type', '')}"
            for r in refs:
                usage.setdefault(fkey(r['table'], r['name']), set()).add(label)

    def slot_list(key):
        return '; '.join(dict.fromkeys(direct.get(key, [])))

    def indirect_list(key):
        own = {x.split(' (')[0] for x in direct.get(key, [])}
        return ', '.join(sorted(set(indirect.get(key, set())) - own))

    out = {'tables': [], 'columns': [], 'measures': [], 'relationships': [],
           'sources': [], 'slots': [], 'missing': []}

    for t in tables:
        ttype = classify_table(t, rels)
        part = next(iter(t['partitions']), {})
        src = part.get('resolvedSource') or {}
        cols_from_origin = sorted({(c.get('origin') or {}).get('object') or '' for c in t['columns']} - {''})
        row = {
            'Nøgle': t['name'], 'Tabel': t['name'], 'Type': ttype,
            'Skjult': 'Ja' if t['isHidden'] else '',
            'Kolonner': len(t['columns']), 'Measures': len(t['measures']),
            'Kilde': source_text(src),
            'Kildeobjekter': ', '.join(cols_from_origin),
            'Opstrøms queries': ', '.join(part.get('references', [])),
            'Query-gruppe': part.get('queryGroup', ''),
            'Lagring': part.get('mode', '') + (' (beregnet)' if t.get('isCalculated') else ''),
            'Dashboard (indirekte)': ', '.join(sorted(indirect.get(t['name'], set()))),
        }
        row.update(merge_editable(t['name'], t.get('description'), ai, prev))
        row['_type'] = 'Tabel'
        out['tables'].append(row)
        if row['Status'] == 'Mangler':
            out['missing'].append({'key': t['name'], 'type': 'Tabel', 'context': {
                'tableType': ttype, 'columns': [c['name'] for c in t['columns']][:40],
                'measures': [m['name'] for m in t['measures']][:40], 'source': row['Kilde'],
                'sourceObjects': row['Kildeobjekter']}})

        pk_cols = {r['to']['column'] for r in rels if r['to']['table'] == t['name']}
        for c in t['columns']:
            if ttype == 'Measures' and c['isHidden']:
                continue
            key = fkey(t['name'], c['name'])
            pii = pii_reason(c['name'], t['name'])
            origin = c.get('origin') or {}
            row = {
                'Nøgle': key, 'Tabel': t['name'], 'Kolonne': c['name'],
                'Datatype': c.get('dataType', ''), 'Format': c.get('formatString', ''),
                'Skjult': 'Ja' if c['isHidden'] else '',
                'Nøglefelt': 'Ja' if (c.get('isKey') or c['name'] in pk_cols) else '',
                'Beregnet': 'DAX' if c.get('isCalculated') else '',
                'Mulig PII': f'Ja ({pii})' if pii else '',
                'Oprindelse': origin_text(origin) if not c.get('isCalculated') else 'Beregnet kolonne (DAX)',
                'Kilde': source_text(origin.get('source')),
                'Power Query-trin': lineage_text(c.get('lineage')),
                'DAX': c.get('expression', '') if c.get('isCalculated') else '',
                'Dashboard-plads': slot_list(key),
                'Dashboard (indirekte)': indirect_list(key),
                'Bruges i rapport': '; '.join(sorted(usage.get(key, set()))),
            }
            row.update(merge_editable(key, c.get('description'), ai, prev,
                                      defaults={'Følsomhed': PII_DEFAULT if pii else ''}))
            if not pii and row['Følsomhed'] == PII_DEFAULT:
                row['Følsomhed'] = ''   # old auto-flag that no longer applies; own text is kept
            row['_type'] = 'Kolonne'
            out['columns'].append(row)
            if row['Status'] == 'Mangler':
                out['missing'].append({'key': key, 'type': 'Kolonne', 'context': {
                    'table': t['name'], 'tableType': ttype, 'dataType': row['Datatype'],
                    'origin': row['Oprindelse'], 'powerQuery': row['Power Query-trin'],
                    'dax': row['DAX'], 'dashboard': row['Dashboard-plads'] or row['Dashboard (indirekte)']}})

        for m in t['measures']:
            key = fkey(t['name'], m['name'])
            d = m.get('dependencies', {})
            cols, tabs, _ = closure(m['name'])
            row = {
                'Nøgle': key, 'Tabel': t['name'], 'Measure': m['name'],
                'Mappe': m.get('displayFolder', ''), 'Format': m.get('formatString', ''),
                'Skjult': 'Ja' if m.get('isHidden') else '',
                'DAX': m.get('expression', ''),
                'Afhænger af measures': ', '.join(d.get('measures', [])),
                'Afhænger af kolonner': ', '.join(fkey(c['table'], c['column']) for c in d.get('columns', [])),
                'Alle kildekolonner': ', '.join(sorted(cols)),
                'Bruges af measures': ', '.join(sorted(used_by.get(m['name'], set()))),
                'Dashboard-plads': slot_list(key),
                'Dashboard (indirekte)': indirect_list(key),
                'Bruges i rapport': '; '.join(sorted(usage.get(key, set()))),
            }
            row.update(merge_editable(key, m.get('description'), ai, prev))
            row['_type'] = 'Measure'
            out['measures'].append(row)
            if row['Status'] == 'Mangler':
                out['missing'].append({'key': key, 'type': 'Measure', 'context': {
                    'dax': m.get('expression', ''), 'format': m.get('formatString', ''),
                    'folder': m.get('displayFolder', ''), 'dependsOn': row['Afhænger af measures'],
                    'sourceColumns': row['Alle kildekolonner'],
                    'dashboard': row['Dashboard-plads'] or row['Dashboard (indirekte)']}})

    for r in rels:
        card = f"{r.get('fromCardinality', 'many')}:{r.get('toCardinality', 'one')}"
        out['relationships'].append({
            'Fra': fkey(r['from']['table'], r['from']['column']),
            'Til': fkey(r['to']['table'], r['to']['column']),
            'Kardinalitet': {'many:one': 'Mange-til-én', 'one:one': 'Én-til-én',
                             'many:many': 'Mange-til-mange', 'one:many': 'Én-til-mange'}.get(card, card),
            'Krydsfilter': 'Begge retninger' if str(r.get('crossFilter', '')).lower().startswith('both') else 'Enkelt',
            'Aktiv': 'Ja' if r.get('active', True) else 'Nej (USERELATIONSHIP)',
        })

    for s in model.get('dataSources', []):
        out['sources'].append({'Type': 'Datakilde', 'Navn': s['kind'], 'Detaljer': s['details'],
                               'Bruges af': ', '.join(s.get('tables', [])), 'Beskrivelse': ''})
    for e in model.get('expressions', []):
        out['sources'].append({
            'Type': 'Parameter' if e['kind'] == 'parameter' else 'Delt query',
            'Navn': e['name'],
            'Detaljer': e.get('parameterValue', '') if e['kind'] == 'parameter' else source_text(e.get('resolvedSource')),
            'Bruges af': ', '.join(e.get('references', [])) if e['kind'] != 'parameter' else '',
            'Query-gruppe': e.get('queryGroup', ''),
            'Beskrivelse': e.get('description', ''),
        })

    lookup = {r['Nøgle']: r for r in out['columns'] + out['measures']}
    for s in slots:
        r = lookup.get(s['field'], {})
        out['slots'].append({
            'Plads': s['slot'], 'Rolle': s['role'], 'Felt': s['field'],
            'Type': r.get('_type', '?'), 'Titel': s.get('label', ''),
            'Beskrivelse': r.get('Beskrivelse', ''),
            'Kildekolonner': r.get('Alle kildekolonner', '') or r.get('Oprindelse', ''),
        })

    current = {r['Nøgle'] for r in out['tables'] + out['columns'] + out['measures']}
    out['removed'] = [{'Nøgle': k, 'Ark': v.get('_sheet', ''), 'Beskrivelse': v.get('Beskrivelse') or '',
                       'Status': v.get('Status') or ''}
                      for k, v in prev.items() if k not in current]
    return out


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def write_excel(cat, path, project, slots_source):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    dark = PatternFill('solid', fgColor='333333')
    teal = PatternFill('solid', fgColor='1FC8C0')
    edit_fill = PatternFill('solid', fgColor='FFF8DC')
    wb = Workbook()
    ws = wb.active
    ws.title = 'Oversigt'

    def status_count(rows):
        c = {s: 0 for s in STATUSES}
        for r in rows:
            c[r['Status']] = c.get(r['Status'], 0) + 1
        return c
    all_rows = cat['tables'] + cat['columns'] + cat['measures']
    sc = status_count(all_rows)
    lines = [
        ('Data catalog', project),
        ('Genereret', datetime.datetime.now().strftime('%Y-%m-%d %H:%M')),
        ('', ''),
        ('Tabeller', len(cat['tables'])),
        ('Kolonner', len(cat['columns'])),
        ('Measures', len(cat['measures'])),
        ('Relationer', len(cat['relationships'])),
        ('Dashboard-felter', f"{len(cat['slots'])} ({slots_source})"),
        ('', ''),
        ('Beskrivelser', ''),
    ] + [(f'  {s}', n) for s, n in sc.items()] + [
        ('', ''),
        ('Fjernet siden sidst', len(cat['removed'])),
        ('', ''),
        ('Sådan redigerer du', 'Ret kun de turkise kolonner (Beskrivelse, Ejer, Følsomhed, Opstrøms kilde, Status).'),
        ('', 'Sæt Status = Godkendt når en beskrivelse må skrives tilbage til modellen (///-beskrivelse i TMDL).'),
        ('', 'Næste kørsel læser dine rettelser og fletter dem ind - de overskrives ikke.'),
        ('Status-værdier', 'Mangler · AI-forslag · Fra model · Redigeret · Godkendt'),
    ]
    for r in lines:
        ws.append(list(r))
    ws['A1'].font = Font(bold=True, size=14)
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 100

    def sheet(name, rows, cols, wide=()):
        s = wb.create_sheet(name)
        s.append(cols)
        for r in rows:
            s.append([r.get(c, '') for c in cols])
        for i, c in enumerate(cols, 1):
            cell = s.cell(row=1, column=i)
            cell.font = Font(bold=True, color='FFFFFF' if c not in EDITABLE else '252423')
            cell.fill = teal if c in EDITABLE else dark
            width = 60 if c in wide else min(45, max(10, len(c) + 2,
                                                       *(len(str(r.get(c, '')))
                                                         for r in rows[:200]) if rows else [10]))
            s.column_dimensions[get_column_letter(i)].width = width
            if c in EDITABLE or c in wide:
                for row in range(2, len(rows) + 2):
                    cell = s.cell(row=row, column=i)
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                    if c in EDITABLE:
                        cell.fill = edit_fill
        if 'Status' in cols and rows:
            col = get_column_letter(cols.index('Status') + 1)
            dv = DataValidation(type='list', formula1='"' + ','.join(STATUSES) + '"', allow_blank=True)
            s.add_data_validation(dv)
            dv.add(f'{col}2:{col}{len(rows) + 1}')
        s.freeze_panes = 'B2' if 'Nøgle' in cols else 'A2'
        if rows:
            s.auto_filter.ref = f'A1:{get_column_letter(len(cols))}{len(rows) + 1}'
        return s

    tail = EDITABLE + ['Beskrivelseskilde']
    sheet('Tabeller', cat['tables'],
          ['Nøgle', 'Type', 'Skjult', 'Kolonner', 'Measures', 'Kilde', 'Kildeobjekter',
           'Opstrøms queries', 'Query-gruppe', 'Lagring', 'Dashboard (indirekte)'] + tail,
          wide=('Beskrivelse',))
    sheet('Kolonner', cat['columns'],
          ['Nøgle', 'Tabel', 'Kolonne', 'Datatype', 'Format', 'Skjult', 'Nøglefelt', 'Beregnet',
           'Mulig PII', 'Oprindelse', 'Kilde', 'Power Query-trin', 'DAX', 'Dashboard-plads',
           'Dashboard (indirekte)', 'Bruges i rapport'] + tail,
          wide=('Beskrivelse', 'Power Query-trin', 'DAX'))
    sheet('Measures', cat['measures'],
          ['Nøgle', 'Tabel', 'Measure', 'Mappe', 'Format', 'Skjult', 'DAX', 'Afhænger af measures',
           'Afhænger af kolonner', 'Alle kildekolonner', 'Bruges af measures', 'Dashboard-plads',
           'Dashboard (indirekte)', 'Bruges i rapport'] + tail,
          wide=('Beskrivelse', 'DAX', 'Alle kildekolonner'))
    sheet('Dashboard', cat['slots'],
          ['Plads', 'Rolle', 'Felt', 'Type', 'Titel', 'Beskrivelse', 'Kildekolonner'],
          wide=('Beskrivelse', 'Kildekolonner'))
    sheet('Relationer', cat['relationships'], ['Fra', 'Til', 'Kardinalitet', 'Krydsfilter', 'Aktiv'])
    sheet('Kilder', cat['sources'], ['Type', 'Navn', 'Detaljer', 'Bruges af', 'Query-gruppe', 'Beskrivelse'],
          wide=('Detaljer', 'Beskrivelse'))
    if cat['removed']:
        sheet('Fjernet', cat['removed'], ['Nøgle', 'Ark', 'Beskrivelse', 'Status'], wide=('Beskrivelse',))
    wb.save(path)


# ---------------------------------------------------------------------------
# Mermaid lineage
# ---------------------------------------------------------------------------

class Mermaid:
    def __init__(self):
        self.ids, self.nodes, self.edges, self.classes = {}, [], [], {}

    def node(self, key, label, cls, shape='box'):
        if key not in self.ids:
            nid = f'n{len(self.ids)}'
            self.ids[key] = nid
            lab = str(label).replace('"', '#quot;').replace('\n', ' ')
            open_, close = {'box': ('["', '"]'), 'round': ('("', '")'),
                            'stadium': ('(["', '"])'), 'hex': ('{{"', '"}}')}[shape]
            self.nodes.append(f'    {nid}{open_}{lab}{close}')
            self.classes.setdefault(cls, []).append(nid)
        return self.ids[key]

    def edge(self, a, b, label=''):
        e = f'    {a} -->|{label}| {b}' if label else f'    {a} --> {b}'
        if e not in self.edges:
            self.edges.append(e)

    def render(self, fence):
        styles = {
            'src': 'fill:#252423,color:#FFFFFF,stroke:#252423',
            'qry': 'fill:#E1DFDD,color:#252423,stroke:#C8C6C4',
            'tbl': 'fill:#333333,color:#FFFFFF,stroke:#333333',
            'col': 'fill:#FFFFFF,color:#252423,stroke:#C8C6C4',
            'msr': 'fill:#1FC8C0,color:#252423,stroke:#1FC8C0',
            'slot': 'fill:#FDEDEF,color:#252423,stroke:#E36C78',
            'page': 'fill:#FDEDEF,color:#252423,stroke:#E36C78',
        }
        body = ['flowchart LR'] + self.nodes + self.edges
        body += [f'    classDef {c} {styles[c]}' for c in self.classes]
        body += [f'    class {",".join(ids)} {c}' for c, ids in self.classes.items()]
        if fence == 'azdo':
            return '::: mermaid\n' + '\n'.join(body) + '\n:::\n'
        return '```mermaid\n' + '\n'.join(body) + '\n```\n'


def write_lineage(parsed, cat, slots, path, fence, project):
    model, report = parsed['model'], parsed['report']
    tables = {t['name']: t for t in model['tables']}
    measure_table = {m['name']: t['name'] for t in model['tables'] for m in t['measures']}
    mdeps = {m['name']: m.get('dependencies', {}) for t in model['tables'] for m in t['measures']}
    exprs = {e['name']: e for e in model.get('expressions', [])}
    out = [f'# Data lineage – {project}', '',
           f"Genereret {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')} af pbip-catalog. "
           'Farver: sort = datakilde, grå = Power Query, mørk = tabel, hvid = kolonne, '
           'turkis = measure, rosa = dashboard.', '']
    if fence == 'azdo':
        out += ['[[_TOC_]]', '']

    # 1. overview
    mm = Mermaid()
    for e in exprs.values():
        if e['kind'] != 'query':
            continue
        q = mm.node('q:' + e['name'], e['name'], 'qry', 'round')
        src = e.get('resolvedSource') or {}
        if src.get('kind') not in (None, 'Other', 'Unknown') and not any(r in exprs and exprs[r]['kind'] == 'query' for r in e.get('references', [])):
            s = mm.node('s:' + source_text(src), source_text(src), 'src', 'stadium')
            mm.edge(s, q)
        for r in e.get('references', []):
            if r in exprs and exprs[r]['kind'] == 'query':
                mm.edge(mm.node('q:' + r, r, 'qry', 'round'), q)
    for t in model['tables']:
        tid = mm.node('t:' + t['name'], t['name'], 'tbl')
        part = next(iter(t['partitions']), {})
        refs = part.get('references', [])
        query_refs = [r for r in refs if r in exprs and exprs[r]['kind'] == 'query']
        table_refs = [r for r in refs if r in tables]
        for r in query_refs:
            mm.edge(mm.node('q:' + r, r, 'qry', 'round'), tid)
        for r in table_refs:
            mm.edge(mm.node('t:' + r, r, 'tbl'), tid)
        if not query_refs and not table_refs:
            src = part.get('resolvedSource') or {}
            if src.get('kind') not in (None, 'Other', 'Unknown'):
                mm.edge(mm.node('s:' + source_text(src), source_text(src), 'src', 'stadium'), tid)
    for page in report.get('pages', []):
        used = set()
        for v in page.get('visuals', []):
            for r in v.get('fieldRefs', []):
                if r['kind'] == 'Measure' and r['name'] in mdeps:
                    used.add(measure_table[r['name']])
                    stack, seen = [r['name']], set()
                    while stack:
                        m = stack.pop()
                        if m in seen:
                            continue
                        seen.add(m)
                        d = mdeps.get(m, {})
                        stack += d.get('measures', [])
                        used |= {c['table'] for c in d.get('columns', [])} | set(d.get('tables', []))
                else:
                    used.add(r['table'])
        if used:
            pid = mm.node('p:' + page['name'], f"Side: {page.get('displayName', page['name'])}", 'page', 'hex')
            for tname in sorted(used):
                if tname in tables:
                    mm.edge(mm.ids.get('t:' + tname) or mm.node('t:' + tname, tname, 'tbl'), pid)
    out += ['## Overblik (tabelniveau)', '',
            'Datakilder → delte Power Query-queries → tabeller → rapportsider.', '', mm.render(fence)]

    # 2. per dashboard slot
    if slots:
        out += ['## Dashboard-pladser', '',
                'Én graf pr. plads på dashboard-siden fra pbip-layout: measure → afhængige measures → '
                'kolonner → oprindelse i kilden.', '']
        col_rows = {r['Nøgle']: r for r in cat['columns']}
        by_slot = {}
        for s in slots:
            by_slot.setdefault(s['slot'], []).append(s)
        for slot, items in by_slot.items():
            mm = Mermaid()
            sid = mm.node('slot', slot, 'slot', 'hex')
            used_cols = set()

            def add_measure(name, parent, seen):
                if name in seen:
                    return mm.ids.get('m:' + name)
                seen.add(name)
                mid = mm.node('m:' + name, f'[{name}]', 'msr', 'round')
                mm.edge(mid, parent)
                d = mdeps.get(name, {})
                for dm in d.get('measures', []):
                    add_measure(dm, mid, seen)
                for c in d.get('columns', []):
                    add_column(fkey(c['table'], c['column']), mid)
                for tb in d.get('tables', []):
                    mm.edge(mm.node('t:' + tb, tb, 'tbl'), mid)
                return mid

            def add_column(key, parent):
                cid = mm.node('c:' + key, key, 'col')
                mm.edge(cid, parent)
                used_cols.add(key)
                r = col_rows.get(key)
                if r and r.get('Oprindelse') and not r['Oprindelse'].startswith('Beregnet'):
                    oid = mm.node('o:' + r['Oprindelse'], r['Oprindelse'], 'src', 'stadium')
                    mm.edge(oid, cid)
            seen = set()
            for it in items:
                t, n = parse_fkey(it['field'])
                if n in mdeps:
                    add_measure(n, sid, seen)
                else:
                    add_column(it['field'], sid)
            out += [f'### {slot}', '', mm.render(fence)]
            rows = [col_rows[k] for k in sorted(used_cols) if k in col_rows]
            if rows:
                out += ['| Kolonne | Oprindelse | Power Query-trin |', '|---|---|---|']
                for r in rows:
                    out.append(f"| `{r['Nøgle']}` | {md(r['Oprindelse'])} | {md(r['Power Query-trin']) or '–'} |")
                out.append('')

    # 3. column lineage per table
    out += ['## Kolonne-lineage pr. tabel', '']
    for t in model['tables']:
        rows = [r for r in cat['columns'] if r['Tabel'] == t['name']]
        if not rows:
            continue
        out += ['<details>', f"<summary><b>{t['name']}</b> ({len(rows)} kolonner)</summary>", '',
                '| Kolonne | Oprindelse | Power Query-trin / DAX |', '|---|---|---|']
        for r in rows:
            out.append(f"| `{r['Kolonne']}` | {md(r['Oprindelse'])} | {md(r['DAX'] or r['Power Query-trin']) or '–'} |")
        out += ['', '</details>', '']
    Path(path).write_text('\n'.join(out), encoding='utf-8')


def md(s):
    return str(s or '').replace('|', '\\|').replace('\n', ' ')


# ---------------------------------------------------------------------------
# DBML (dbdiagram.io / dbdocs.io)
# ---------------------------------------------------------------------------

def dq(name):
    return '"' + str(name).replace('\\', '\\\\').replace('"', '\\"') + '"'


def sq(text):
    return "'" + re.sub(r'\s+', ' ', str(text or '')).strip().replace('\\', '\\\\').replace("'", "\\'") + "'"


def tq(text):
    return "'''" + str(text or '').replace('\\', '\\\\').replace("'''", "\\'''") + "'''"


GROUP_COLORS = {'Fakta': '#333333', 'Dimension': '#605E5C', 'Dato': '#252423',
                'Bro': '#605E5C', 'Measures': '#1FC8C0', 'Hjælpe': '#C8C6C4'}


def write_dbml(parsed, cat, path, project):
    model = parsed['model']
    rels = model['relationships']
    trows = {r['Nøgle']: r for r in cat['tables']}
    crows = {r['Nøgle']: r for r in cat['columns']}
    mrows = {r['Nøgle']: r for r in cat['measures']}
    out = [f'// Genereret af pbip-catalog {datetime.date.today().isoformat()} – rediger beskrivelser i catalog.xlsx',
           f'Project {dq(project)} {{', "  database_type: 'Power BI'",
           '  Note: ' + tq(f'Power BI semantisk model "{project}". Tabeller, kolonner og measures med '
                           'beskrivelser, oprindelse og dashboard-pladser i noterne.'), '}', '']
    groups = {}
    for t in model['tables']:
        tr = trows[t['name']]
        ttype = tr['Type']
        groups.setdefault(ttype, []).append(t['name'])
        pk = next((c['name'] for c in t['columns'] if c.get('isKey')), None) or \
            next((r['to']['column'] for r in rels if r['to']['table'] == t['name']), None)
        out.append(f"Table {dq(t['name'])} [headercolor: {GROUP_COLORS.get(ttype, '#C8C6C4')}] {{")
        for c in t['columns']:
            row = crows.get(fkey(t['name'], c['name']))
            if not row:
                continue
            settings = []
            if c['name'] == pk:
                settings.append('pk')
            note = [row['Beskrivelse']]
            if c['isHidden']:
                note.append('Skjult')
            if c.get('isCalculated'):
                note.append('Beregnet (DAX): ' + short(c.get('expression', ''), 200))
            elif row['Oprindelse']:
                note.append('Oprindelse: ' + row['Oprindelse'])
            if row['Mulig PII']:
                note.append('Mulig PII')
            if row['Dashboard-plads']:
                note.append('Dashboard: ' + row['Dashboard-plads'])
            note = ' · '.join(x for x in note if x)
            if note:
                settings.append('note: ' + sq(note))
            dtype = re.sub(r'\W', '', c.get('dataType') or 'unknown') or 'unknown'
            out.append(f"  {dq(c['name'])} {dtype}" + (f" [{', '.join(settings)}]" if settings else ''))
        for m in t['measures']:
            row = mrows[fkey(t['name'], m['name'])]
            note = [row['Beskrivelse'], 'DAX: ' + short(m.get('expression', ''), 300)]
            if row['Dashboard-plads']:
                note.append('Dashboard: ' + row['Dashboard-plads'])
            elif row['Dashboard (indirekte)']:
                note.append('Bruges indirekte i: ' + row['Dashboard (indirekte)'])
            out.append(f"  {dq(m['name'])} measure [note: {sq(' · '.join(x for x in note if x))}]")
        tnote = [tr['Beskrivelse'], f'Type: {ttype}']
        if tr['Kilde']:
            tnote.append(f"Kilde: {tr['Kilde']}")
        if tr['Kildeobjekter']:
            tnote.append(f"Kildeobjekter: {tr['Kildeobjekter']}")
        if tr['Opstrøms queries']:
            tnote.append(f"Power Query: {tr['Opstrøms queries']}")
        inactive = [f"{r['from']['column']} → {r['to']['table']}.{r['to']['column']}"
                    for r in rels if r['from']['table'] == t['name'] and not r.get('active', True)]
        if inactive:
            tnote.append('Inaktive relationer (USERELATIONSHIP): ' + '; '.join(inactive))
        both = [f"{r['from']['column']} ↔ {r['to']['table']}.{r['to']['column']}"
                for r in rels if r['from']['table'] == t['name']
                and str(r.get('crossFilter', '')).lower().startswith('both')]
        if both:
            tnote.append('Krydsfilter begge retninger: ' + '; '.join(both))
        if tr['Dashboard (indirekte)']:
            tnote.append(f"Dashboard: {tr['Dashboard (indirekte)']}")
        if not t['columns'] and not t['measures']:
            out.append('  "(ingen kolonner)" unknown')
        out += ['', '  Note: ' + tq('\n'.join(x for x in tnote if x)), '}', '']
    for g in ('Fakta', 'Bro', 'Dimension', 'Dato', 'Measures', 'Hjælpe'):
        if groups.get(g):
            out.append(f'TableGroup {dq(g)} {{')
            out += [f'  {dq(t)}' for t in groups[g]]
            out += ['}', '']
    for r in rels:
        op = {('many', 'one'): '>', ('one', 'many'): '<', ('one', 'one'): '-',
              ('many', 'many'): '<>'}.get((r.get('fromCardinality', 'many'), r.get('toCardinality', 'one')), '>')
        flags = []
        if not r.get('active', True):
            flags.append('inaktiv (USERELATIONSHIP)')
        if str(r.get('crossFilter', '')).lower().startswith('both'):
            flags.append('krydsfilter begge retninger')
        comment = f"  // {', '.join(flags)}" if flags else ''
        out.append(f"Ref: {dq(r['from']['table'])}.{dq(r['from']['column'])} {op} "
                   f"{dq(r['to']['table'])}.{dq(r['to']['column'])}{comment}")
    Path(path).write_text('\n'.join(out) + '\n', encoding='utf-8')


# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('pbip')
    ap.add_argument('--out')
    ap.add_argument('--layout-config')
    ap.add_argument('--previous')
    ap.add_argument('--ai-descriptions')
    ap.add_argument('--fence', choices=['azdo', 'github'], default='azdo')
    ap.add_argument('--parser')
    a = ap.parse_args()

    pbip = Path(a.pbip).resolve()
    out = Path(a.out).resolve() if a.out else pbip / 'docs'
    out.mkdir(parents=True, exist_ok=True)
    parser, parser_path = load_parser(a.parser)
    parsed = {'model': parser.parse_model(pbip), 'report': parser.parse_report(pbip)}
    project = parsed['model']['name'] or pbip.name

    cfg_path = Path(a.layout_config) if a.layout_config else out / 'layout-config.json'
    if cfg_path.is_file():
        slots = slots_from_config(json.loads(cfg_path.read_text(encoding='utf-8')))
        slots_source = f'fra {cfg_path.name}'
    else:
        page = next((p for p in parsed['report'].get('pages', []) if p['name'] == LAYOUT_PAGE_ID), None)
        slots = slots_from_page(page) if page else []
        slots_source = 'udledt af dashboard-siden' if page else 'ingen pbip-layout-side fundet'

    xlsx = out / 'catalog.xlsx'
    prev_path = Path(a.previous) if a.previous else xlsx
    prev = read_previous(prev_path)
    ai = json.loads(Path(a.ai_descriptions).read_text(encoding='utf-8')) if a.ai_descriptions else {}

    pii_reason = getattr(parser, 'pii_reason', None) or \
        (lambda col, table: 'navnemønster' if parser.is_pii(col) else '')
    cat = build(parsed, slots, prev, ai, pii_reason)
    if xlsx.is_file():
        shutil.copy2(xlsx, out / 'catalog.backup.xlsx')
    write_excel(cat, xlsx, project, slots_source)
    write_lineage(parsed, cat, slots, out / 'lineage.md', a.fence, project)
    write_dbml(parsed, cat, out / 'model.dbml', project)
    strip = lambda rows: [{k: v for k, v in r.items() if not k.startswith('_')} for r in rows]  # noqa: E731
    (out / 'catalog.json').write_text(json.dumps({
        'project': project, 'generated': datetime.datetime.now().isoformat(timespec='seconds'),
        'slotsSource': slots_source,
        **{k: strip(v) if isinstance(v, list) else v for k, v in cat.items() if k != 'missing'},
    }, indent=2, ensure_ascii=False), encoding='utf-8')
    (out / 'missing_descriptions.json').write_text(
        json.dumps(cat['missing'], indent=2, ensure_ascii=False), encoding='utf-8')

    rows = cat['tables'] + cat['columns'] + cat['measures']
    counts = {s: sum(1 for r in rows if r['Status'] == s) for s in STATUSES}
    print(json.dumps({
        'project': project, 'out': str(out), 'parser': parser_path,
        'tables': len(cat['tables']), 'columns': len(cat['columns']), 'measures': len(cat['measures']),
        'dashboardFields': len(cat['slots']), 'slotsSource': slots_source,
        'descriptionStatus': counts, 'missingDescriptions': len(cat['missing']),
        'removedSinceLastRun': len(cat['removed']),
    }, indent=2, ensure_ascii=False))


if __name__ == '__main__':
    main()
